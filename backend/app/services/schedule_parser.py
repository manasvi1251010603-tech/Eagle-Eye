import os
from typing import Dict, Any, List
import openpyxl
from app.config import BASE_DIR


LOGICAL_FIELDS = {
    "device_ref": ["device", "component", "ref", "source_device", "item"],
    "source_terminal": ["source_terminal", "from_terminal", "src_term", "source_term", "terminal_1"],
    "dest_device": ["dest_device", "destination_device", "to_device", "target_device"],
    "dest_terminal": ["dest_terminal", "destination_terminal", "to_terminal", "dst_term", "terminal_2"],
    "ferrule_id": ["ferrule", "ferrule_id", "wire_tag", "tag", "marker", "sleeve"],
    "wire_colour": ["color", "colour", "wire_color", "wire_colour", "insulation"],
    "wire_gauge": ["gauge", "size", "cross_section", "mm2", "wire_gauge_mm2"]
}


def parse_schedule_excel(excel_relative_or_abs_path: str, selected_sheet: str = None) -> Dict[str, Any]:
    """
    Parses an Excel wiring schedule, returns sheets, columns, preview rows,
    and automatic column mappings with validation warnings.
    """
    if os.path.isabs(excel_relative_or_abs_path):
        full_path = excel_relative_or_abs_path
    else:
        full_path = os.path.join(BASE_DIR, excel_relative_or_abs_path)

    if not os.path.exists(full_path):
        # Fallback to demo workbook
        full_path = os.path.join(BASE_DIR, "data", "demo", "Wiring_Schedule_RevC.xlsx")

    wb = openpyxl.load_workbook(full_path, data_only=True)
    sheet_names = wb.sheetnames

    if not selected_sheet or selected_sheet not in sheet_names:
        selected_sheet = "Wiring_Schedule" if "Wiring_Schedule" in sheet_names else sheet_names[0]

    ws = wb[selected_sheet]

    rows_iter = ws.iter_rows(values_only=True)
    raw_headers = next(rows_iter, None)

    if not raw_headers:
        return {
            "sheets": sheet_names,
            "selected_sheet": selected_sheet,
            "columns": [],
            "sample_rows": [],
            "total_rows": 0,
            "detected_mappings": {},
            "validation_warnings": ["The selected worksheet is completely empty."]
        }

    columns = [str(h).strip() if h is not None else f"Column_{idx+1}" for idx, h in enumerate(raw_headers)]

    sample_rows = []
    total_rows = 0
    for r in rows_iter:
        if any(cell is not None for cell in r):
            total_rows += 1
            if len(sample_rows) < 15:
                row_dict = {}
                for idx, col_name in enumerate(columns):
                    val = r[idx] if idx < len(r) else None
                    row_dict[col_name] = str(val) if val is not None else ""
                sample_rows.append(row_dict)

    # Intelligent Auto-Detection of logical column mappings
    detected_mappings = {}
    validation_warnings = []

    for logical_key, candidates in LOGICAL_FIELDS.items():
        matched_col = None
        for col in columns:
            col_lower = col.lower().replace(" ", "_").replace("-", "_")
            if any(cand in col_lower for cand in candidates):
                matched_col = col
                break
        if matched_col:
            detected_mappings[logical_key] = matched_col
        else:
            if logical_key in ["device_ref", "ferrule_id", "wire_colour"]:
                validation_warnings.append(
                    f"Required field '{logical_key}' could not be automatically mapped. Please map it manually."
                )

    wb.close()

    return {
        "sheets": sheet_names,
        "selected_sheet": selected_sheet,
        "columns": columns,
        "sample_rows": sample_rows,
        "total_rows": total_rows,
        "detected_mappings": detected_mappings,
        "validation_warnings": validation_warnings
    }
