import os
import json
import uuid
import datetime
from typing import List, Optional, Dict, Any
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Query
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app.config import BASE_DIR, UPLOADS_DIR, REPORTS_DIR, DEMO_DATA_DIR
from app.models.models import InspectionSession, InspectionFinding, RepairAction, AuditEvent, GeneratedReport
from app.schemas.schemas import (
    InspectionCreate, InspectionSummary, InspectionDetail, FindingResponse,
    FindingReviewUpdate, RepairActionResponse, RepairActionCreate, RepairActionUpdate,
    AlignmentRequest, AlignmentResponse, ColumnMappingRequest, SchedulePreviewResponse,
    ReinspectionRequest, ReinspectionDiffResponse, ReportResponse
)
from app.services.file_validation import save_uploaded_file, ALLOWED_IMAGE_EXTS, ALLOWED_PDF_EXTS, ALLOWED_EXCEL_EXTS
from app.services.ga_parser import parse_ga_pdf
from app.services.schedule_parser import parse_schedule_excel
from app.services.alignment import compute_perspective_alignment
from app.services.comparison import execute_inspection_pipeline
from app.services.reinspection import compute_reinspection_diff
from app.services.report_generator import generate_pdf_report, generate_excel_report

router = APIRouter(prefix="/inspections", tags=["Inspections"])


@router.post("", response_model=InspectionSummary)
def create_inspection_session(payload: InspectionCreate, db: Session = Depends(get_db)):
    session_id = f"INS-{datetime.datetime.utcnow().strftime('%Y%m%d')}-{uuid.uuid4().hex[:4].upper()}"

    session = InspectionSession(
        id=session_id,
        cabinet_id=payload.cabinet_id,
        work_order_id=payload.work_order_id,
        batch_id=payload.batch_id,
        inspector=payload.inspector,
        reference_revision=payload.reference_revision or "Rev C",
        status="DRAFT",
        processing_mode=payload.processing_mode or "DEMO",
        cabinet_image_path="data/demo/cabinet_sample_01.jpg",
        ga_filename="GA_Drawing_Cabinet_RevC.pdf",
        schedule_filename="Wiring_Schedule_RevC.xlsx"
    )
    db.add(session)

    # Audit event
    db.add(AuditEvent(
        id=f"AUD-{uuid.uuid4().hex[:8]}",
        session_id=session.id,
        event_type="SESSION_CREATED",
        description=f"Inspection session {session.id} initialized for Cabinet {session.cabinet_id}.",
        actor=session.inspector
    ))
    db.commit()
    db.refresh(session)
    return _build_session_summary(session, db)


@router.get("", response_model=List[InspectionSummary])
def list_inspection_sessions(
    status: Optional[str] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db)
):
    query = db.query(InspectionSession).order_by(InspectionSession.created_at.desc())

    if status and status != "ALL":
        query = query.filter(InspectionSession.status == status)
    if search:
        search_fmt = f"%{search}%"
        query = query.filter(
            (InspectionSession.cabinet_id.like(search_fmt)) |
            (InspectionSession.id.like(search_fmt)) |
            (InspectionSession.work_order_id.like(search_fmt)) |
            (InspectionSession.inspector.like(search_fmt))
        )

    sessions = query.all()
    return [_build_session_summary(s, db) for s in sessions]


@router.get("/{session_id}", response_model=InspectionDetail)
def get_inspection_detail(session_id: str, db: Session = Depends(get_db)):
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    summary = _build_session_summary(session, db)
    findings = db.query(InspectionFinding).filter(InspectionFinding.session_id == session_id).all()
    audits = db.query(AuditEvent).filter(AuditEvent.session_id == session_id).order_by(AuditEvent.created_at.desc()).all()

    return InspectionDetail(
        **summary.model_dump(),
        column_mapping=session.column_mapping,
        alignment_data=session.alignment_data,
        findings=[_build_finding_response(f) for f in findings],
        audit_events=audits
    )


@router.post("/{session_id}/upload-files")
async def upload_inspection_files(
    session_id: str,
    cabinet_photo: Optional[UploadFile] = File(None),
    ga_drawing: Optional[UploadFile] = File(None),
    wiring_schedule: Optional[UploadFile] = File(None),
    db: Session = Depends(get_db)
):
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    uploaded = []
    if cabinet_photo:
        orig_name, rel_path = await save_uploaded_file(cabinet_photo, ALLOWED_IMAGE_EXTS)
        session.cabinet_image_path = rel_path
        uploaded.append(f"Cabinet Photo: {orig_name}")

    if ga_drawing:
        orig_name, rel_path = await save_uploaded_file(ga_drawing, ALLOWED_PDF_EXTS)
        session.ga_filename = orig_name
        session.ga_file_path = rel_path
        uploaded.append(f"GA Drawing: {orig_name}")

    if wiring_schedule:
        orig_name, rel_path = await save_uploaded_file(wiring_schedule, ALLOWED_EXCEL_EXTS)
        session.schedule_filename = orig_name
        session.schedule_file_path = rel_path
        uploaded.append(f"Wiring Schedule: {orig_name}")

    if uploaded:
        db.add(AuditEvent(
            id=f"AUD-{uuid.uuid4().hex[:8]}",
            session_id=session.id,
            event_type="FILES_UPLOADED",
            description=f"Uploaded {', '.join(uploaded)}.",
            actor=session.inspector
        ))
        db.commit()

    return {
        "message": "Files uploaded successfully",
        "cabinet_image_path": session.cabinet_image_path,
        "ga_filename": session.ga_filename,
        "schedule_filename": session.schedule_filename
    }


@router.post("/{session_id}/load-demo-pack")
def load_demo_pack(session_id: str, db: Session = Depends(get_db)):
    """
    Instantly equips the session with verified sample industrial cabinet files.
    """
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    session.cabinet_image_path = "data/demo/cabinet_sample_01.jpg"
    session.repaired_image_path = "data/demo/cabinet_sample_01_repaired.jpg"
    session.ga_filename = "GA_Drawing_Cabinet_RevC.pdf"
    session.schedule_filename = "Wiring_Schedule_RevC.xlsx"
    session.processing_mode = "DEMO"

    db.add(AuditEvent(
        id=f"AUD-{uuid.uuid4().hex[:8]}",
        session_id=session.id,
        event_type="DEMO_PACK_LOADED",
        description="Pre-loaded realistic industrial reference pack (Cabinet CAB-4092, GA Rev C, Schedule Rev C).",
        actor=session.inspector
    ))
    db.commit()
    return {"message": "Demo sample pack loaded successfully into inspection workspace"}


@router.post("/{session_id}/preview-schedule", response_model=SchedulePreviewResponse)
def preview_schedule(
    session_id: str,
    selected_sheet: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    sched_path = session.schedule_file_path or "data/demo/Wiring_Schedule_RevC.xlsx"
    return parse_schedule_excel(sched_path, selected_sheet=selected_sheet)


@router.post("/{session_id}/map-columns")
def save_column_mapping(
    session_id: str,
    payload: ColumnMappingRequest,
    db: Session = Depends(get_db)
):
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    session.column_mapping = json.dumps({
        "sheet_name": payload.sheet_name,
        "mappings": payload.mappings
    })
    db.commit()
    return {"message": "Column mappings saved successfully"}


@router.post("/{session_id}/preview-ga")
def preview_ga(session_id: str, db: Session = Depends(get_db)):
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    ga_path = session.ga_file_path or "data/demo/GA_Drawing_Cabinet_RevC.pdf"
    return parse_ga_pdf(ga_path)


@router.post("/{session_id}/align", response_model=AlignmentResponse)
def align_images(
    session_id: str,
    payload: AlignmentRequest,
    db: Session = Depends(get_db)
):
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    photo_path = session.cabinet_image_path or "data/demo/cabinet_sample_01.jpg"
    cad_path = "uploads/ga_preview_page_1.png"

    pts_photo = [p.model_dump() for p in payload.points_photo] if payload.points_photo else None
    pts_cad = [p.model_dump() for p in payload.points_cad] if payload.points_cad else None

    scale = payload.scale if payload.scale is not None else 1.0
    trans_x = payload.trans_x if payload.trans_x is not None else 0.0
    trans_y = payload.trans_y if payload.trans_y is not None else 0.0
    rotation = payload.rotation if payload.rotation is not None else 0.0
    opacity = payload.opacity if payload.opacity is not None else 0.45

    result = compute_perspective_alignment(
        photo_rel_path=photo_path,
        cad_rel_path=cad_path,
        manual_points_photo=pts_photo,
        manual_points_cad=pts_cad,
        scale=scale,
        trans_x=trans_x,
        trans_y=trans_y,
        rotation=rotation,
        opacity=opacity
    )

    session.alignment_data = json.dumps(result)
    db.add(AuditEvent(
        id=f"AUD-{uuid.uuid4().hex[:8]}",
        session_id=session.id,
        event_type="ALIGNMENT_EXECUTED",
        description=f"Perspective alignment executed: {result['alignment_status']} (scale: {scale}, rot: {rotation} deg, offset: [{trans_x}, {trans_y}]).",
        actor=session.inspector
    ))
    db.commit()

    return AlignmentResponse(
        session_id=session.id,
        alignment_status=result["alignment_status"],
        confidence_score=result["confidence_score"],
        scale=result.get("scale", scale),
        trans_x=result.get("trans_x", trans_x),
        trans_y=result.get("trans_y", trans_y),
        rotation=result.get("rotation", rotation),
        opacity=result.get("opacity", opacity),
        homography_matrix=result.get("homography_matrix"),
        warning_message=result.get("warning_message"),
        aligned_overlay_url=result.get("aligned_overlay_url")
    )


@router.get("/{session_id}/extracted-references")
def get_extracted_references(session_id: str, db: Session = Depends(get_db)):
    """
    Returns structured components, terminal connections, and GA layout references.
    If the user has manually edited them, returns the saved version.
    """
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    if session.extracted_references:
        try:
            return json.loads(session.extracted_references)
        except Exception:
            pass

    # Extract default structured references from files
    sched_path = session.schedule_file_path or "data/demo/Wiring_Schedule_RevC.xlsx"
    full_sched = os.path.join(BASE_DIR, sched_path) if not os.path.isabs(sched_path) else sched_path

    components = []
    wires = []

    try:
        import openpyxl
        wb = openpyxl.load_workbook(full_sched, data_only=True)
        if "BOM_List" in wb.sheetnames:
            ws_bom = wb["BOM_List"]
            for r in ws_bom.iter_rows(min_row=2, values_only=True):
                if r and r[0]:
                    components.append({
                        "ref": str(r[0]),
                        "type": str(r[1]) if len(r) > 1 and r[1] else "",
                        "manufacturer": str(r[2]) if len(r) > 2 and r[2] else "",
                        "part_number": str(r[3]) if len(r) > 3 and r[3] else "",
                        "rail": str(r[4]) if len(r) > 4 and r[4] else "DIN 1",
                        "qty": int(r[5]) if len(r) > 5 and r[5] else 1
                    })
        if "Wiring_Schedule" in wb.sheetnames:
            ws_wire = wb["Wiring_Schedule"]
            for r in ws_wire.iter_rows(min_row=2, values_only=True):
                if r and r[0]:
                    wires.append({
                        "wire_tag": str(r[0]),
                        "source_device": str(r[1]) if len(r) > 1 and r[1] else "",
                        "source_terminal": str(r[2]) if len(r) > 2 and r[2] else "",
                        "dest_device": str(r[3]) if len(r) > 3 and r[3] else "",
                        "dest_terminal": str(r[4]) if len(r) > 4 and r[4] else "",
                        "expected_color": str(r[5]) if len(r) > 5 and r[5] else "Black",
                        "wire_gauge": str(r[6]) if len(r) > 6 and r[6] else "1.5",
                        "ferrule_id": str(r[7]) if len(r) > 7 and r[7] else str(r[0]),
                        "signal_type": str(r[8]) if len(r) > 8 and r[8] else "Control"
                    })
        wb.close()
    except Exception as e:
        print("Excel extraction fallback:", e)

    # Fallback default components if empty
    if not components:
        components = [
            {"ref": "PS1", "type": "Power Supply 24VDC 10A", "manufacturer": "Phoenix Contact", "part_number": "TRIO-PS-2G", "rail": "DIN 1", "qty": 1},
            {"ref": "QF1", "type": "Main Breaker 3P 63A", "manufacturer": "Schneider Electric", "part_number": "A9F74363", "rail": "DIN 1", "qty": 1},
            {"ref": "QF2", "type": "Branch Breaker 1P 16A", "manufacturer": "Schneider Electric", "part_number": "A9F74116", "rail": "DIN 1", "qty": 1},
            {"ref": "QF3", "type": "Branch Breaker 1P 10A", "manufacturer": "Schneider Electric", "part_number": "A9F74110", "rail": "DIN 1", "qty": 1},
            {"ref": "SPD1", "type": "Surge Protector 40kA", "manufacturer": "DEHN", "part_number": "952070", "rail": "DIN 1", "qty": 1},
            {"ref": "PLC1", "type": "PLC CPU 1214C", "manufacturer": "Siemens", "part_number": "6ES7214-1AG40-0XB0", "rail": "DIN 2", "qty": 1},
            {"ref": "K1", "type": "Motor Contactor 25A", "manufacturer": "Siemens", "part_number": "3RT2026-1BB40", "rail": "DIN 2", "qty": 1},
            {"ref": "K2", "type": "Motor Contactor 25A", "manufacturer": "Siemens", "part_number": "3RT2026-1BB40", "rail": "DIN 2", "qty": 1},
            {"ref": "KA1", "type": "Safety Relay Module", "manufacturer": "Pilz", "part_number": "PNOZ s4", "rail": "DIN 2", "qty": 1},
            {"ref": "X1", "type": "Terminal Block Strip 20P", "manufacturer": "Weidmuller", "part_number": "WDU 2.5", "rail": "DIN 3", "qty": 20}
        ]

    if not wires:
        wires = [
            {"wire_tag": "W-101", "source_device": "QF1", "source_terminal": "L1_OUT", "dest_device": "QF2", "dest_terminal": "L_IN", "expected_color": "Black", "wire_gauge": "2.5", "ferrule_id": "QF1:L1", "signal_type": "AC Power"},
            {"wire_tag": "W-105", "source_device": "PS1", "source_terminal": "+24V", "dest_device": "PLC1", "dest_terminal": "L+", "expected_color": "Red", "wire_gauge": "1.5", "ferrule_id": "PLC1:24V", "signal_type": "DC 24V"},
            {"wire_tag": "W-107", "source_device": "PLC1", "source_terminal": "DO_0", "dest_device": "K1", "dest_terminal": "A1", "expected_color": "Dark Blue", "wire_gauge": "0.75", "ferrule_id": "K1:A1", "signal_type": "Control"},
            {"wire_tag": "W-108", "source_device": "K1", "source_terminal": "A2", "dest_device": "X1", "dest_terminal": "14", "expected_color": "Blue", "wire_gauge": "0.75", "ferrule_id": "X1:14", "signal_type": "Control"},
            {"wire_tag": "W-110", "source_device": "K2", "source_terminal": "A2", "dest_device": "X1", "dest_terminal": "15", "expected_color": "Blue", "wire_gauge": "0.75", "ferrule_id": "X1:15", "signal_type": "Control"},
            {"wire_tag": "W-113", "source_device": "QF1", "source_terminal": "PE", "dest_device": "X1", "dest_terminal": "6", "expected_color": "Green-Yellow", "wire_gauge": "4.0", "ferrule_id": "PE:X1-6", "signal_type": "Ground"}
        ]

    ref_payload = {
        "drawing_number": session.ga_filename or "GA-CAB-4092-C.pdf",
        "revision": session.reference_revision or "Rev C",
        "enclosure_dimensions": "1000 x 800 x 300 mm (RAL 7035)",
        "identified_rails": ["DIN 1 (Top Power)", "DIN 2 (Middle Control)", "DIN 3 (Bottom Field Terminals)"],
        "components": components,
        "wires": wires,
        "is_user_edited": False
    }

    return ref_payload


@router.post("/{session_id}/extracted-references")
def update_extracted_references(
    session_id: str,
    payload: Dict[str, Any],
    db: Session = Depends(get_db)
):
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    payload["is_user_edited"] = True
    session.extracted_references = json.dumps(payload)
    db.add(AuditEvent(
        id=f"AUD-{uuid.uuid4().hex[:8]}",
        session_id=session.id,
        event_type="REFERENCES_MANUALLY_EDITED",
        description=f"Inspector {session.inspector} edited and verified parsed reference tables.",
        actor=session.inspector
    ))
    db.commit()
    return {"message": "Extracted references updated and saved successfully"}


@router.post("/{session_id}/run-pipeline")
def run_inspection_pipeline(
    session_id: str,
    is_repaired: bool = False,
    db: Session = Depends(get_db)
):
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    session.status = "PROCESSING"
    db.commit()

    # Clear existing findings for this session if re-running
    db.query(InspectionFinding).filter(InspectionFinding.session_id == session_id).delete()

    image_path = session.repaired_image_path if is_repaired else (session.cabinet_image_path or "data/demo/cabinet_sample_01.jpg")

    findings_data = execute_inspection_pipeline(
        session_id=session.id,
        image_rel_path=image_path,
        ga_data={},
        schedule_data={},
        is_repaired=is_repaired
    )

    fail_count = 0
    review_count = 0

    for item in findings_data:
        finding = InspectionFinding(
            id=item["id"],
            session_id=session.id,
            check_type=item["check_type"],
            component_ref=item["component_ref"],
            expected_value=item["expected_value"],
            observed_value=item["observed_value"],
            status=item["status"],
            severity=item.get("severity", "LOW" if item["status"] == "PASS" else ("HIGH" if item["status"] == "FAIL" else "MEDIUM")),
            reason=item["reason"],
            recommended_action=item.get("recommended_action"),
            confidence=item["confidence"],
            is_simulated=item.get("is_simulated", True),
            bbox_x=item["bbox_x"],
            bbox_y=item["bbox_y"],
            bbox_width=item["bbox_width"],
            bbox_height=item["bbox_height"],
            evidence_crop_path=item["evidence_crop_path"],
            source_reference=item["source_reference"],
            repair_status=item["repair_status"]
        )
        db.add(finding)
        db.flush()

        if item["status"] == "FAIL":
            fail_count += 1
            # Add guided repair actions
            if item["check_type"] == "GA_PLACEMENT":
                actions = [
                    "Verify lock-out / tag-out (LOTO) isolation on cabinet.",
                    f"Unclip {item['component_ref']} from current DIN rail mounting slot.",
                    f"Shift and mount {item['component_ref']} to correct designated position.",
                    "Re-tighten clamps and apply torque seal inspection paint."
                ]
                for act in actions:
                    db.add(RepairAction(
                        id=f"REP-{uuid.uuid4().hex[:8]}",
                        finding_id=finding.id,
                        checklist_item=act,
                        is_completed=False,
                        technician=session.inspector
                    ))
            elif item["check_type"] == "FERRULE_TEXT":
                actions = [
                    f"Isolate wiring circuit for {item['component_ref']}.",
                    "Remove incorrectly printed ferrule tag.",
                    f"Print corrected ferrule sleeve '{item['expected_value']}'.",
                    "Re-terminate conductor and perform mechanical pull tug test."
                ]
                for act in actions:
                    db.add(RepairAction(
                        id=f"REP-{uuid.uuid4().hex[:8]}",
                        finding_id=finding.id,
                        checklist_item=act,
                        is_completed=False,
                        technician=session.inspector
                    ))

        if item["status"] == "REVIEW":
            review_count += 1

    if fail_count > 0:
        session.status = "REPAIR_REQUIRED"
    elif review_count > 0:
        session.status = "NEEDS_REVIEW"
    else:
        session.status = "READY_FOR_REVIEW"

    db.add(AuditEvent(
        id=f"AUD-{uuid.uuid4().hex[:8]}",
        session_id=session.id,
        event_type="INSPECTION_RUN",
        description=f"AI Inspection pipeline completed: {len(findings_data)} checked, {fail_count} FAIL, {review_count} REVIEW.",
        actor="EAGLE EYE Engine"
    ))
    db.commit()

    return {
        "message": "Inspection pipeline completed",
        "session_status": session.status,
        "total_findings": len(findings_data),
        "fail_count": fail_count,
        "review_count": review_count
    }


@router.get("/{session_id}/findings", response_model=List[FindingResponse])
def get_session_findings(
    session_id: str,
    status: Optional[str] = None,
    check_type: Optional[str] = None,
    search: Optional[str] = None,
    db: Session = Depends(get_db)
):
    query = db.query(InspectionFinding).filter(InspectionFinding.session_id == session_id)

    if status and status != "ALL":
        query = query.filter(InspectionFinding.status == status)
    if check_type and check_type != "ALL":
        query = query.filter(InspectionFinding.check_type == check_type)
    if search:
        search_fmt = f"%{search}%"
        query = query.filter(
            (InspectionFinding.component_ref.like(search_fmt)) |
            (InspectionFinding.id.like(search_fmt)) |
            (InspectionFinding.expected_value.like(search_fmt)) |
            (InspectionFinding.observed_value.like(search_fmt)) |
            (InspectionFinding.reason.like(search_fmt))
        )

    findings = query.all()
    return [_build_finding_response(f) for f in findings]


@router.patch("/findings/{finding_id}/review", response_model=FindingResponse)
def update_finding_review(
    finding_id: str,
    payload: FindingReviewUpdate,
    db: Session = Depends(get_db)
):
    finding = db.query(InspectionFinding).filter(InspectionFinding.id == finding_id).first()
    if not finding:
        raise HTTPException(status_code=404, detail="Finding not found")

    finding.reviewer_status = payload.reviewer_status
    finding.reviewer_notes = payload.reviewer_notes
    if payload.status:
        finding.status = payload.status

    db.add(AuditEvent(
        id=f"AUD-{uuid.uuid4().hex[:8]}",
        session_id=finding.session_id,
        finding_id=finding.id,
        event_type="FINDING_OVERRIDDEN",
        description=f"Reviewer override on {finding.component_ref}: status set to {finding.status} ({payload.reviewer_status}). Note: {payload.reviewer_notes}",
        actor="Human QC Reviewer"
    ))
    db.commit()
    db.refresh(finding)
    return _build_finding_response(finding)


@router.post("/findings/{finding_id}/repair-actions", response_model=RepairActionResponse)
def add_repair_action(
    finding_id: str,
    payload: RepairActionCreate,
    db: Session = Depends(get_db)
):
    finding = db.query(InspectionFinding).filter(InspectionFinding.id == finding_id).first()
    if not finding:
        raise HTTPException(status_code=404, detail="Finding not found")

    action = RepairAction(
        id=f"REP-{uuid.uuid4().hex[:8]}",
        finding_id=finding_id,
        checklist_item=payload.checklist_item,
        is_completed=payload.is_completed,
        technician=payload.technician,
        notes=payload.notes
    )
    db.add(action)
    db.commit()
    db.refresh(action)
    return action


@router.patch("/repair-actions/{action_id}", response_model=RepairActionResponse)
def update_repair_action(
    action_id: str,
    payload: RepairActionUpdate,
    db: Session = Depends(get_db)
):
    action = db.query(RepairAction).filter(RepairAction.id == action_id).first()
    if not action:
        raise HTTPException(status_code=404, detail="Repair action not found")

    action.is_completed = payload.is_completed
    if payload.notes is not None:
        action.notes = payload.notes
    if payload.technician is not None:
        action.technician = payload.technician

    # Check finding's overall repair status
    finding = action.finding
    all_actions = db.query(RepairAction).filter(RepairAction.finding_id == finding.id).all()
    if all(a.is_completed for a in all_actions):
        finding.repair_status = "REPAIRED"
    else:
        finding.repair_status = "IN_PROGRESS"

    db.add(AuditEvent(
        id=f"AUD-{uuid.uuid4().hex[:8]}",
        session_id=finding.session_id,
        finding_id=finding.id,
        event_type="REPAIR_CHECKLIST_UPDATED",
        description=f"Checklist item '{action.checklist_item[:40]}...' marked as {'COMPLETED' if action.is_completed else 'PENDING'}.",
        actor=action.technician or "Technician"
    ))
    db.commit()
    db.refresh(action)
    return action


@router.post("/{session_id}/reinspect", response_model=ReinspectionDiffResponse)
def create_reinspection(
    session_id: str,
    payload: ReinspectionRequest,
    db: Session = Depends(get_db)
):
    """
    Creates a new versioned inspection session linked to parent_session_id.
    Executes inspection on the post-repair image and computes the diff.
    """
    parent = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not parent:
        raise HTTPException(status_code=404, detail="Parent inspection session not found")

    new_session_id = f"INS-{datetime.datetime.utcnow().strftime('%Y%m%d')}-R{uuid.uuid4().hex[:3].upper()}"

    new_session = InspectionSession(
        id=new_session_id,
        cabinet_id=parent.cabinet_id,
        work_order_id=parent.work_order_id,
        batch_id=parent.batch_id,
        inspector=payload.inspector,
        reference_revision=parent.reference_revision,
        status="PROCESSING",
        parent_session_id=parent.id,
        cabinet_image_path="data/demo/cabinet_sample_01_repaired.jpg",
        repaired_image_path="data/demo/cabinet_sample_01_repaired.jpg",
        ga_filename=parent.ga_filename,
        schedule_filename=parent.schedule_filename,
        processing_mode="DEMO"
    )
    db.add(new_session)
    db.commit()

    # Run inspection pipeline for repaired state
    findings_data = execute_inspection_pipeline(
        session_id=new_session.id,
        image_rel_path=new_session.cabinet_image_path,
        ga_data={},
        schedule_data={},
        is_repaired=True
    )

    for item in findings_data:
        db.add(InspectionFinding(
            id=item["id"],
            session_id=new_session.id,
            check_type=item["check_type"],
            component_ref=item["component_ref"],
            expected_value=item["expected_value"],
            observed_value=item["observed_value"],
            status=item["status"],
            severity=item.get("severity", "LOW" if item["status"] == "PASS" else ("HIGH" if item["status"] == "FAIL" else "MEDIUM")),
            reason=item["reason"],
            recommended_action=item.get("recommended_action"),
            confidence=item["confidence"],
            is_simulated=item.get("is_simulated", True),
            bbox_x=item["bbox_x"],
            bbox_y=item["bbox_y"],
            bbox_width=item["bbox_width"],
            bbox_height=item["bbox_height"],
            evidence_crop_path=item["evidence_crop_path"],
            source_reference=item["source_reference"],
            repair_status="VERIFIED" if item["status"] == "PASS" else "PENDING"
        ))
    db.commit()

    parent_findings = db.query(InspectionFinding).filter(InspectionFinding.session_id == parent.id).all()
    new_findings = db.query(InspectionFinding).filter(InspectionFinding.session_id == new_session.id).all()

    diff_result = compute_reinspection_diff(parent_findings, new_findings)

    new_session.status = "PASSED" if diff_result["can_close_session"] else "REPAIR_REQUIRED"
    db.add(AuditEvent(
        id=f"AUD-{uuid.uuid4().hex[:8]}",
        session_id=new_session.id,
        event_type="REINSPECTION_RUN",
        description=f"Reinspection session {new_session.id} linked to {parent.id}. Resolved: {len(diff_result['resolved_findings'])}, Persistent: {len(diff_result['persistent_findings'])}.",
        actor=payload.inspector
    ))
    db.commit()

    return ReinspectionDiffResponse(
        reinspection_session_id=new_session.id,
        parent_session_id=parent.id,
        resolved_findings=[_build_finding_response(f) for f in diff_result["resolved_findings"]],
        persistent_findings=[_build_finding_response(f) for f in diff_result["persistent_findings"]],
        new_findings=[_build_finding_response(f) for f in diff_result["new_findings"]],
        review_needed_findings=[_build_finding_response(f) for f in diff_result["review_needed_findings"]],
        can_close_session=diff_result["can_close_session"],
        closure_blockers=diff_result["closure_blockers"]
    )


@router.post("/{session_id}/close")
def close_session(session_id: str, db: Session = Depends(get_db)):
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    findings = db.query(InspectionFinding).filter(InspectionFinding.session_id == session_id).all()
    blocking_fails = [f for f in findings if f.status == "FAIL"]
    blocking_reviews = [f for f in findings if f.status == "REVIEW" and not f.reviewer_status]

    if blocking_fails:
        raise HTTPException(
            status_code=400,
            detail=f"Cannot close session with {len(blocking_fails)} unresolved FAIL defects. Rework and reinspection required."
        )
    if blocking_reviews:
        raise HTTPException(
            status_code=400,
            detail=f"Cannot close session with {len(blocking_reviews)} unreviewed items. Human reviewer disposition required."
        )

    session.status = "CLOSED"
    db.add(AuditEvent(
        id=f"AUD-{uuid.uuid4().hex[:8]}",
        session_id=session.id,
        event_type="SESSION_CLOSED",
        description="Inspection session officially CLOSED and archived into QC audit repository.",
        actor=session.inspector
    ))
    db.commit()
    return {"message": "Inspection session closed successfully", "status": session.status}


@router.post("/{session_id}/reports/pdf", response_model=ReportResponse)
def export_pdf_report(session_id: str, db: Session = Depends(get_db)):
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    findings = db.query(InspectionFinding).filter(InspectionFinding.session_id == session_id).all()
    pdf_rel_path = generate_pdf_report(session, findings)

    report_rec = GeneratedReport(
        id=f"RPT-{uuid.uuid4().hex[:8]}",
        session_id=session.id,
        report_type="PDF",
        file_path=pdf_rel_path,
        filename=os.path.basename(pdf_rel_path)
    )
    db.add(report_rec)
    db.add(AuditEvent(
        id=f"AUD-{uuid.uuid4().hex[:8]}",
        session_id=session.id,
        event_type="REPORT_GENERATED",
        description=f"Generated PDF QC Report: {report_rec.filename}.",
        actor="System"
    ))
    db.commit()
    db.refresh(report_rec)

    return ReportResponse(
        id=report_rec.id,
        session_id=report_rec.session_id,
        report_type=report_rec.report_type,
        file_path=report_rec.file_path,
        filename=report_rec.filename,
        download_url=f"/reports/{report_rec.filename}",
        created_at=report_rec.created_at
    )


@router.post("/{session_id}/reports/excel", response_model=ReportResponse)
def export_excel_report(session_id: str, db: Session = Depends(get_db)):
    session = db.query(InspectionSession).filter(InspectionSession.id == session_id).first()
    if not session:
        raise HTTPException(status_code=404, detail="Inspection session not found")

    findings = db.query(InspectionFinding).filter(InspectionFinding.session_id == session_id).all()
    excel_rel_path = generate_excel_report(session, findings)

    report_rec = GeneratedReport(
        id=f"RPT-{uuid.uuid4().hex[:8]}",
        session_id=session.id,
        report_type="EXCEL",
        file_path=excel_rel_path,
        filename=os.path.basename(excel_rel_path)
    )
    db.add(report_rec)
    db.add(AuditEvent(
        id=f"AUD-{uuid.uuid4().hex[:8]}",
        session_id=session.id,
        event_type="REPORT_GENERATED",
        description=f"Generated Excel QC Report: {report_rec.filename}.",
        actor="System"
    ))
    db.commit()
    db.refresh(report_rec)

    return ReportResponse(
        id=report_rec.id,
        session_id=report_rec.session_id,
        report_type=report_rec.report_type,
        file_path=report_rec.file_path,
        filename=report_rec.filename,
        download_url=f"/reports/{report_rec.filename}",
        created_at=report_rec.created_at
    )


def _build_session_summary(session: InspectionSession, db: Session) -> InspectionSummary:
    findings = db.query(InspectionFinding).filter(InspectionFinding.session_id == session.id).all()
    p_cnt = sum(1 for f in findings if f.status == "PASS")
    f_cnt = sum(1 for f in findings if f.status == "FAIL")
    r_cnt = sum(1 for f in findings if f.status == "REVIEW")
    nc_cnt = sum(1 for f in findings if f.status == "NOT_CHECKED")

    return InspectionSummary(
        id=session.id,
        cabinet_id=session.cabinet_id,
        work_order_id=session.work_order_id,
        batch_id=session.batch_id,
        inspector=session.inspector,
        reference_revision=session.reference_revision or "Rev C",
        status=session.status,
        parent_session_id=session.parent_session_id,
        processing_mode=session.processing_mode or "DEMO",
        cabinet_image_path=session.cabinet_image_path,
        repaired_image_path=session.repaired_image_path,
        ga_filename=session.ga_filename,
        schedule_filename=session.schedule_filename,
        pass_count=p_cnt,
        fail_count=f_cnt,
        review_count=r_cnt,
        not_checked_count=nc_cnt,
        total_findings=len(findings),
        created_at=session.created_at,
        updated_at=session.updated_at
    )


def _build_finding_response(f: InspectionFinding) -> FindingResponse:
    repairs = [
        RepairActionResponse(
            id=r.id,
            finding_id=r.finding_id,
            checklist_item=r.checklist_item,
            is_completed=r.is_completed,
            technician=r.technician,
            notes=r.notes,
            updated_at=r.updated_at
        ) for r in f.repair_actions
    ]

    return FindingResponse(
        id=f.id,
        session_id=f.session_id,
        original_finding_id=f.original_finding_id,
        check_type=f.check_type,
        component_ref=f.component_ref,
        expected_value=f.expected_value,
        observed_value=f.observed_value,
        status=f.status,
        severity=f.severity or "MEDIUM",
        reason=f.reason,
        recommended_action=f.recommended_action,
        confidence=f.confidence,
        is_simulated=f.is_simulated if f.is_simulated is not None else True,
        bbox_x=f.bbox_x,
        bbox_y=f.bbox_y,
        bbox_width=f.bbox_width,
        bbox_height=f.bbox_height,
        evidence_crop_path=f.evidence_crop_path,
        source_reference=f.source_reference,
        reviewer_status=f.reviewer_status,
        reviewer_notes=f.reviewer_notes,
        repair_status=f.repair_status,
        repair_actions=repairs,
        created_at=f.created_at
    )
